import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.extractor import extract_text, extract_text_from_txt_or_md


def test_extract_txt_decodifica_e_remove_espacos_nas_pontas():
    file_bytes = "  FastAPI usa type hints.  ".encode("utf-8")

    result = extract_text_from_txt_or_md(file_bytes)

    assert result == [{"page_number": 1, "text": "FastAPI usa type hints."}]


def test_extract_txt_ignora_bytes_invalidos_em_vez_de_quebrar():
    # 0xFF sozinho não é um caractere UTF-8 válido.
    file_bytes = b"texto valido \xff depois"

    result = extract_text_from_txt_or_md(file_bytes)

    assert "texto valido" in result[0]["text"]


def test_extract_docx_junta_paragrafos_nao_vazios():
    doc = DocxDocument()
    doc.add_paragraph("Primeiro parágrafo.")
    doc.add_paragraph("")  # parágrafo vazio deve ser ignorado
    doc.add_paragraph("Segundo parágrafo.")
    buffer = io.BytesIO()
    doc.save(buffer)

    result = extract_text(buffer.getvalue(), "docx")

    assert result == [
        {"page_number": 1, "text": "Primeiro parágrafo.\nSegundo parágrafo."}
    ]


def test_extract_pdf_pula_paginas_sem_texto():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)  # página sem nenhum texto
    buffer = io.BytesIO()
    writer.write(buffer)

    result = extract_text(buffer.getvalue(), "pdf")

    assert result == []


def test_dispatcher_delega_pela_extensao_sem_case_sensitivity():
    file_bytes = "conteúdo".encode("utf-8")

    result_minuscula = extract_text(file_bytes, "txt")
    result_com_ponto_maiuscula = extract_text(file_bytes, ".TXT")

    assert result_minuscula == result_com_ponto_maiuscula


def test_dispatcher_rejeita_formato_nao_suportado():
    with pytest.raises(ValueError, match="não suportado"):
        extract_text(b"dados quaisquer", "exe")
